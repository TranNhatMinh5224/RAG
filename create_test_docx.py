import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    """Set background color for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_document():
    doc = docx.Document()

    # Configure Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Document Header / Agency
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_header.add_run("TẬP ĐOÀN CÔNG NGHỆ & ĐỔI MỚI SÁNG TẠO NEXUS CORP\n")
    r1.bold = True
    r1.font.size = Pt(13)
    r1.font.color.rgb = RGBColor(30, 41, 59)
    
    r2 = p_header.add_run("TRUNG TÂM GIẢI PHÁP ĐÁM MÂY & TRÍ TUỆ NHÂN TẠO (CLOUD & AI EXCELLENCE)\n")
    r2.bold = True
    r2.font.size = Pt(11)
    r2.font.color.rgb = RGBColor(71, 85, 105)
    
    r3 = p_header.add_run("Số: 108/2026/QĐ-NEXUS-AWS • Văn phòng: Tầng 7, 36 Cát Linh, Đống Đa, Hà Nội\n")
    r3.italic = True
    r3.font.size = Pt(10)
    r3.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph("─" * 45).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("QUY CHẾ QUẢN TRỊ HẠ TẦNG ĐÁM MÂY AWS\nVÀ VẬN HÀNH TRÍ TUỆ NHÂN TẠO AMAZON BEDROCK 2026")
    r_title.bold = True
    r_title.font.size = Pt(16)
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run("(Ban hành kèm theo Quyết định số 108/QĐ-TGĐ ngày 15 tháng 08 năm 2026)")
    r_meta.italic = True
    r_meta.font.size = Pt(10.5)

    doc.add_paragraph()

    # --- CHƯƠNG I ---
    h1 = doc.add_heading("CHƯƠNG I: QUY ĐỊNH CHUNG", level=1)
    
    doc.add_heading("Điều 1: Phạm vi điều chỉnh và Đối tượng áp dụng", level=2)
    p = doc.add_paragraph()
    p.add_run("1. Khoản 1: ").bold = True
    p.add_run("Quy chế này quy định toàn diện các tiêu chuẩn kỹ thuật, quy trình bảo mật mạng Zero-Trust, chính sách sao lưu dữ liệu và định mức chi phí áp dụng cho toàn bộ hạ tầng đám mây Amazon Web Services (AWS) và nền tảng Trí tuệ Nhân tạo thế hệ mới Amazon Bedrock.")

    p = doc.add_paragraph()
    p.add_run("2. Khoản 2: ").bold = True
    p.add_run("Quy chế này có hiệu lực bắt buộc áp dụng đối với toàn thể Kỹ sư Điện toán đám mây (Cloud Engineer), Kỹ sư Trí tuệ nhân tạo (AI Engineer), Chuyên viên Vận hành DevOps, Thực tập sinh Công nghệ và các bên thứ ba tham gia phát triển dự án NexusDoc AI tại Trung tâm.")

    doc.add_heading("Điều 2: Giải thích thuật ngữ và Định nghĩa kỹ thuật", level=2)
    p = doc.add_paragraph()
    p.add_run("Trong văn bản này, các thuật ngữ dưới đây được định nghĩa thống nhất như sau:\n")
    p.add_run("a) Điểm a: ").bold = True
    p.add_run("'Hệ thống RAG (Retrieval-Augmented Generation)' là kiến trúc kết hợp giữa tìm kiếm tài liệu nội bộ và mô hình ngôn ngữ lớn để trả lời có căn cứ trích dẫn nguồn xác thực.\n")
    p.add_run("b) Điểm b: ").bold = True
    p.add_run("'Amazon Bedrock Mantle Endpoint' là cổng giao tiếp dịch vụ AI serverless thế hệ mới của AWS đặt tại khu vực US East (N. Virginia), cung cấp các mô hình nền tảng như Mistral 14B, Amazon Nova và Qwen.\n")
    p.add_run("c) Điểm c: ").bold = True
    p.add_run("'Mô hình Zero-Trust' là nguyên tắc kiến trúc bảo mật 'Không bao giờ tin tưởng, luôn luôn xác thực', cô lập hoàn toàn cơ sở dữ liệu trong mạng riêng nội bộ.")

    # --- CHƯƠNG II ---
    doc.add_heading("CHƯƠNG II: TIÊU CHUẨN AN TOÀN HẠ TẦNG ĐÁM MÂY AWS", level=1)

    doc.add_heading("Điều 3: Phân vùng mạng VPC và Kiểm soát cổng truy cập", level=2)
    p = doc.add_paragraph()
    p.add_run("1. Khoản 1: ").bold = True
    p.add_run("Hạ tầng VPC của NexusDoc AI bắt buộc quy hoạch theo dải mạng 10.0.0.0/16 trải dài trên tối thiểu 2 Availability Zones (ap-southeast-1a và ap-southeast-1b), chia thành 3 lớp phân cấp:\n")
    p.add_run("   - Public Subnet: Chỉ tiếp nhận lưu lượng Internet qua Application Load Balancer (ALB).\n")
    p.add_run("   - Private Subnet: Đặt cụm container tính toán ECS Fargate Backend và Celery Worker.\n")
    p.add_run("   - Isolated Subnet: Lưu trữ cơ sở dữ liệu quan hệ Amazon RDS PostgreSQL và Qdrant Vector Engine, tuyệt đối không gắn Internet Gateway.")

    p = doc.add_paragraph()
    p.add_run("2. Khoản 2: ").bold = True
    p.add_run("Nghiêm cấm mở cổng Internet công cộng trực tiếp cho cổng 5432 (PostgreSQL) và cổng 6333 (Qdrant). Mọi truy cập quản trị phải thông qua AWS Systems Manager Session Manager hoặc VPN nội bộ.")

    doc.add_heading("Điều 4: Quản lý chứng thư và Khóa bí mật (AWS Secrets Manager)", level=2)
    p = doc.add_paragraph()
    p.add_run("1. Khoản 1: ").bold = True
    p.add_run("Toàn bộ các khóa bí mật bao gồm chuỗi kết nối Database, JWT SECRET_KEY, GEMINI_API_KEY và BEDROCK_API_KEY bắt buộc phải lưu trữ tập trung tại dịch vụ AWS Secrets Manager với tên định danh 'rag/production/credentials'.")

    p = doc.add_paragraph()
    p.add_run("2. Khoản 2: ").bold = True
    p.add_run("Nghiêm cấm hành vi đính kèm (hard-code) các khóa API trong mã nguồn đẩy lên GitHub. Bất kỳ cá nhân nào vi phạm quy định này sẽ bị xử lý nghiêm khắc theo ")
    p.add_run("Khoản 2 Điều 10 của Quy chế này").bold = True
    p.add_run(".")

    # --- CHƯƠNG III ---
    doc.add_heading("CHƯƠNG III: QUẢN TRỊ TRÍ TUỆ NHÂN TẠO & AMAZON BEDROCK", level=1)

    doc.add_heading("Điều 5: Danh mục Mô hình Ngôn ngữ được phê duyệt", level=2)
    p = doc.add_paragraph()
    p.add_run("1. Khoản 1: ").bold = True
    p.add_run("Hệ thống NexusDoc AI chỉ được phép kích hoạt các Foundation Models đã qua thẩm định an toàn tại Bedrock Mantle Console:\n")
    p.add_run("   - Mô hình chính thức cấp doanh nghiệp: 'mistral.ministral-3-14b-instruct'.\n")
    p.add_run("   - Mô hình suy luận tốc độ cao: 'amazon.nova-micro-v1:0'.\n")
    p.add_run("   - Mô hình fallback môi trường Dev: Google Gemini 2.5 Flash.")

    p = doc.add_paragraph()
    p.add_run("2. Khoản 2: ").bold = True
    p.add_run("Chính sách đa nhà cung cấp (Multi-Provider): Khi cờ cấu hình USE_BEDROCK=true, hệ thống định tuyến 100% truy vấn sang Amazon Bedrock Mantle qua đường truyền bảo mật TLS 1.3.")

    doc.add_heading("Điều 6: Cơ chế An toàn Bedrock Guardrails và Chống ảo giác", level=2)
    p = doc.add_paragraph()
    p.add_run("1. Khoản 1 (Bảo vệ thông tin cá nhân - PII Redaction): ").bold = True
    p.add_run("Bộ lọc Guardrails tự động kiểm duyệt và che dấu các trường dữ liệu nhạy cảm bao gồm: Số Căn cước công dân (CCCD 12 chữ số), Số tài khoản ngân hàng, Mã số thuế doanh nghiệp và Mật khẩu truy cập.")

    p = doc.add_paragraph()
    p.add_run("2. Khoản 2 (Nguyên tắc Strict Grounding): ").bold = True
    p.add_run("AI chỉ được phép sinh câu trả lời khi có bằng chứng trực tiếp trong tài liệu nội bộ trích xuất. Nếu ngữ cảnh không có thông tin, AI bắt buộc phải trả lời câu chuẩn: 'Tài liệu nội bộ hiện tại không đề cập đến nội dung này'.")

    # --- CHƯƠNG IV ---
    doc.add_heading("CHƯƠNG IV: CHỈ TIÊU SLA & THANG ĐO XỬ LÝ SỰ CỐ", level=1)

    doc.add_heading("Điều 7: Cam kết chất lượng dịch vụ kỹ thuật (SLA 2026)", level=2)
    p = doc.add_paragraph()
    p.add_run("Toàn bộ các chỉ tiêu kỹ thuật đo lường thực tế trên AWS CloudWatch và Application Load Balancer phải tuân thủ nghiêm ngặt các ngưỡng cam kết sau:")

    # Table 1: SLA Metrics
    table1 = doc.add_table(rows=1, cols=4)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    table1.style = 'Table Grid'
    hdr_cells = table1.rows[0].cells
    hdr_cells[0].text = "Chỉ số Kỹ thuật (Metric)"
    hdr_cells[1].text = "Ngưỡng Cam kết (SLA)"
    hdr_cells[2].text = "Kết quả Thực nghiệm"
    hdr_cells[3].text = "Đánh giá Vận hành"
    
    for c in hdr_cells:
        set_cell_background(c, "1E293B")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

    sla_data = [
        ("Độ trễ toàn trình (ALB Latency p95)", "≤ 3.0 giây", "1.82 giây", "Đạt chuẩn xuất sắc"),
        ("Độ trễ truy vấn Vector Qdrant (p99)", "≤ 50.0 mili-giây", "14.2 mili-giây", "Vượt chỉ tiêu"),
        ("Thời gian phục hồi sự cố (MTTR)", "≤ 180 giây", "65 giây", "Tự phục hồi nhanh"),
        ("Độ chính xác trích xuất (Precision@3)", "≥ 85.0%", "94.2%", "Đạt chuẩn cao"),
        ("Tỷ lệ sẵn sàng hệ sinh thái (Uptime)", "≥ 99.9%", "99.95%", "Chuẩn AWS Enterprise"),
        ("Tỷ lệ chặn ảo giác (Zero-Hallucination)", "100%", "100%", "Đạt tuyệt đối")
    ]

    for row_idx, data in enumerate(sla_data):
        row_cells = table1.add_row().cells
        bg_color = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for i, text in enumerate(data):
            row_cells[i].text = text
            set_cell_background(row_cells[i], bg_color)
            if i == 2:
                row_cells[i].paragraphs[0].runs[0].font.bold = True

    doc.add_paragraph()

    doc.add_heading("Điều 8: Phân loại cấp độ sự cố kỹ thuật (Incident Severity)", level=2)
    p = doc.add_paragraph()
    p.add_run("Khi phát sinh sự cố vận hành hạ tầng đám mây và mô hình AI, đội ngũ trực vận hành (On-call Team) phải tuân thủ phân cấp thời gian phản hồi quy định tại Bảng 2:")

    # Table 2: Severity Matrix
    table2 = doc.add_table(rows=1, cols=4)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    table2.style = 'Table Grid'
    hdr_cells2 = table2.rows[0].cells
    hdr_cells2[0].text = "Cấp độ Sự cố"
    hdr_cells2[1].text = "Mô tả Tác động Nghiệp vụ"
    hdr_cells2[2].text = "Thời gian Phản hồi (Target)"
    hdr_cells2[3].text = "Chế độ Báo cáo"

    for c in hdr_cells2:
        set_cell_background(c, "0F766E")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

    severity_data = [
        ("Cấp P1 (Khẩn cấp)", "Toàn bộ cụm ALB hoặc Bedrock mất kết nối, người dùng không thể tra cứu tài liệu.", "Dưới 15 phút", "Kích hoạt SNS báo động TGĐ & Trưởng nhóm Cloud"),
        ("Cấp P2 (Nghiêm trọng)", "Một phân hệ gặp lỗi (ví dụ Qdrant Vector DB bị nghẽn CPU > 85%, hoặc ElastiCache failover).", "Dưới 1 giờ", "Báo cáo Kỹ sư trưởng qua kênh Slack On-call"),
        ("Cấp P3 (Trung bình)", "Tốc độ phản hồi bị suy giảm nhẹ (> 2.5s), các chức năng cốt lõi vẫn vận hành bình thường.", "Dưới 4 giờ", "Ghi nhận vé xử lý Jira và điều chỉnh Auto-Scaling"),
        ("Cấp P4 (Thấp)", "Yêu cầu hỗ trợ thông tin, đổi cấu hình prompt, cập nhật tài liệu kiểm thử mới.", "Dưới 24 giờ", "Xử lý trong giờ làm việc hành chính")
    ]

    for row_idx, data in enumerate(severity_data):
        row_cells = table2.add_row().cells
        bg_color = "F0FDFA" if row_idx % 2 == 0 else "FFFFFF"
        for i, text in enumerate(data):
            row_cells[i].text = text
            set_cell_background(row_cells[i], bg_color)
            if i == 0:
                row_cells[i].paragraphs[0].runs[0].font.bold = True

    doc.add_paragraph()

    doc.add_heading("Điều 9: Chế tài xử lý vi phạm an toàn thông tin", level=2)
    p = doc.add_paragraph()
    p.add_run("1. Khoản 1: ").bold = True
    p.add_run("Mọi hành vi cố tình vô hiệu hóa Bedrock Guardrails, tiết lộ mã bí mật BEDROCK_API_KEY hoặc làm rò rỉ tài liệu nhạy cảm của doanh nghiệp ra ngoài hệ sinh thái AWS sẽ bị xử lý nghiêm khắc theo chế tài quy định tại ")
    p.add_run("Khoản 2 Điều 10 của Quy chế này").bold = True
    p.add_run(".")

    doc.add_heading("Điều 10: Biện pháp Kỷ luật và Trách nhiệm Bồi thường", level=2)
    p = doc.add_paragraph()
    p.add_run("1. Khoản 1: ").bold = True
    p.add_run("Cá nhân vi phạm có nghĩa vụ bồi hoàn toàn bộ tổn thất tài chính phát sinh do chi phí tài nguyên điện toán đám mây vượt mức hoặc chi phí rò rỉ dữ liệu.")

    p = doc.add_paragraph()
    p.add_run("2. Khoản 2 (Các hình thức kỷ luật cụ thể): ").bold = True
    p.add_run("Tùy theo tính chất và mức độ vi phạm, người vi phạm sẽ chịu một trong các hình thức kỷ luật sau:\n")
    p.add_run("   - Mức 1: Khiển trách bằng văn bản và tước quyền truy cập AWS Console trong thời hạn 14 ngày đối với vi phạm lần đầu không gây hậu quả nghiêm trọng.\n")
    p.add_run("   - Mức 2: Thu hồi toàn bộ quyền IAM Role, đình chỉ công tác 30 ngày đối với hành vi làm lộ mã bí mật API Key.\n")
    p.add_run("   - Mức 3: Buộc thôi việc và chuyển hồ sơ sang cơ quan pháp luật có thẩm quyền xử lý theo Luật An ninh mạng đối với hành vi cố ý phá hoại cơ sở dữ liệu.")

    # Signatures
    doc.add_paragraph()
    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_sign1 = p_sign.add_run("Hà Nội, ngày 15 tháng 08 năm 2026\n")
    r_sign1.italic = True
    r_sign2 = p_sign.add_run("TỔNG GIÁM ĐỐC TẬP ĐOÀN NEXUS CORP\n")
    r_sign2.bold = True
    r_sign3 = p_sign.add_run("(Đã ký và đóng dấu)\n\n\n")
    r_sign4 = p_sign.add_run("TS. NGUYỄN VĂN AN")
    r_sign4.bold = True

    output_path = "c:/Users/Minh/Desktop/TTTN/RAG/Quy_Che_Van_Hanh_Ha_Tang_Cloud_Va_AI_Bedrock_2026.docx"
    doc.save(output_path)
    print(f"Document created successfully at: {output_path}")

if __name__ == "__main__":
    create_document()

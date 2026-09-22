import fitz  # PyMuPDF
import docx
import pandas as pd
from pptx import Presentation
import re
import os
import numpy as np
try:
    from paddleocr import PaddleOCR
except ImportError:
    PaddleOCR = None

from langchain_text_splitters import RecursiveCharacterTextSplitter

try:
    from services.legal_parser import LegalDocumentParser
except ImportError:
    from src.backend.services.legal_parser import LegalDocumentParser

class DocumentProcessor:
    def __init__(self, embeddings=None):
        """Khởi tạo với Recursive Splitter (cắt theo cấu trúc)"""
        self.recursive_splitter = RecursiveCharacterTextSplitter(
            chunk_size=1200,
            chunk_overlap=150,
            separators=["\n\n", "\n", ". ", " ", ""]
        )
        self._ocr = None

    @property
    def ocr(self):
        """Lazy load PaddleOCR chỉ khi thật sự gặp trang/ảnh cần OCR"""
        if self._ocr is None and PaddleOCR is not None:
            try:
                print("[INFO] Dang khoi tao PaddleOCR (Lazy Load)...")
                self._ocr = PaddleOCR(use_angle_cls=True, lang='vi', show_log=False)
            except Exception as e:
                print(f"[WARN] Khong the nap PaddleOCR ({e})")
                self._ocr = None
        return self._ocr

    def clean_text(self, text: str) -> str:
        # 1. Nối lại các từ bị gãy ở cuối dòng do dấu gạch nối (Ví dụ: "trách nhi- \n ệm")
        text = re.sub(r'(\w+)-\s+(\w+)', r'\1\2', text)
        
        # 2. Xóa các ký tự ASCII điều khiển (không xóa ký tự tiếng Việt)
        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
        
        # 3. Gom các khoảng trắng, dấu tab, dấu xuống dòng liên tiếp thành 1 dấu cách
        text = re.sub(r'\s+', ' ', text)
        
        return text.strip()

    @staticmethod
    def is_toc_page(text: str) -> bool:
        """Phát hiện xem trang có phải là trang Mục lục (Table of Contents) không"""
        clean = text.strip()
        if not clean:
            return False
        has_toc_header = bool(re.search(r'^\s*(?:Table\s+of\s+Contents|Contents|Mục\s+lục)\b', clean, re.IGNORECASE | re.MULTILINE))
        dot_leaders = len(re.findall(r'(?:\.\s*){3,}', clean))
        lines = [l.strip() for l in clean.split('\n') if l.strip()]
        toc_lines = sum(1 for l in lines if re.search(r'[\.\-\s]{2,}\s*\d+$', l) or re.search(r'^(?:chapter|\d+\.|\bappendix\b).*\s+\d+$', l, re.IGNORECASE))
        if has_toc_header and (dot_leaders >= 2 or toc_lines >= 3):
            return True
        if dot_leaders >= 4 or (len(lines) >= 5 and toc_lines / len(lines) >= 0.4):
            return True
        return False

    HEADING_PATTERNS = [
        r'^(?:chapter|chương|phần|mục)\s+[\dIVXLCDM]+[:\.\s\-]+.*$',
        r'^\d+(?:\.\d+)*\s+[A-Z\u00C0-\u1EF9].*$',
        r'^[A-Z\u00C0-\u1EF9\s\:\-]{4,60}$',
        r'^appendix\s+[A-Z\d][:.\s\-]+.*$',
        r'^Điều\s+\d+[\.:]?\s*.*$'
    ]

    @classmethod
    def extract_heading(cls, line: str) -> str | None:
        stripped = line.strip()
        if not stripped or len(stripped) > 100 or len(stripped) < 3:
            return None
        if stripped.endswith((';', '...', ',')):
            return None
        for pattern in cls.HEADING_PATTERNS:
            if re.match(pattern, stripped, re.IGNORECASE):
                return stripped
        return None

    @staticmethod
    def classify_chunk_type(heading: str, text: str) -> str:
        h = (heading or "").lower()
        t = text[:200].lower()
        combined = f"{h} {t}"
        
        if any(k in combined for k in ["abstract", "tóm tắt"]):
            return "abstract"
        if any(k in combined for k in ["introduction", "motivation", "giới thiệu", "động lực", "bối cảnh", "problem statement", "objective"]):
            return "introduction"
        if any(k in combined for k in ["methodology", "architecture", "method", "phương pháp", "kiến trúc", "mô hình", "framework"]):
            return "methodology"
        if any(k in combined for k in ["experiment", "result", "thực nghiệm", "thử nghiệm", "kết quả", "evaluation", "benchmark", "auc"]):
            return "experiment"
        if any(k in combined for k in ["conclusion", "kết luận", "discussion", "thảo luận", "future work"]):
            return "conclusion"
        return "general"

    @staticmethod
    def is_junk_chunk(text: str) -> bool:
        clean = text.strip()
        if len(clean) < 35:
            return True
        letters = re.findall(r'[a-zA-Z0-9\u00C0-\u1EF9]', clean)
        if not letters or len(letters) / len(clean) < 0.35:
            return True
        if re.search(r'(?:\.\s*){4,}', clean):
            return True
        if re.search(r'\bcontents\b', clean, re.IGNORECASE) and re.search(r'\babstract\b', clean, re.IGNORECASE) and re.search(r'\d+\s+\d+', clean):
            return True
        return False

    def process_pdf(self, file_path: str):
        """Hàm phụ: Đọc file PDF (Hỗ trợ Text-PDF, Scan-PDF bằng OCR và Lọc bỏ Mục lục)"""
        doc = fitz.open(file_path)
        pages_text = []
        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            text = page.get_text("text").strip()
            if not text and self.ocr is not None:
                try:
                    pix = page.get_pixmap()
                    img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.h, pix.w, pix.n)
                    if pix.n == 4:
                        import cv2
                        img = cv2.cvtColor(img, cv2.COLOR_RGBA2RGB)
                    elif pix.n == 1:
                        import cv2
                        img = cv2.cvtColor(img, cv2.COLOR_GRAY2RGB)
                    
                    result = self.ocr.ocr(img, cls=True)
                    page_text = ""
                    if result and result[0]:
                        for line in result[0]:
                            page_text += line[1][0] + "\n"
                    text = page_text.strip()
                except Exception as e:
                    print(f"[ERROR] Loi khi OCR trang {page_num + 1}: {e}")
                    
            if text:
                if self.is_toc_page(text):
                    print(f"[INFO] [Cleaner] Bo qua trang Muc luc (TOC): Trang {page_num + 1}")
                    continue
                pages_text.append({"text": text, "page": page_num + 1})
                
        doc.close()
        return pages_text

    def process_image(self, file_path: str):
        """Hàm phụ: Đọc trực tiếp file ảnh bằng PaddleOCR"""
        if self.ocr is None:
            return []
        try:
            result = self.ocr.ocr(file_path, cls=True)
            page_text = ""
            if result and result[0]:
                for line in result[0]:
                    page_text += line[1][0] + "\n"
            return [{"text": page_text, "page": 1}] if page_text.strip() else []
        except Exception as e:
            print(f"[ERROR] Loi khi OCR anh {file_path}: {e}")
            return []

    def process_docx(self, file_path: str):
        """Hàm phụ: Đọc file Word (.docx)"""
        doc = docx.Document(file_path)
        full_text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
        return [{"text": full_text, "page": 1}] if full_text else []

    def process_xlsx(self, file_path: str):
        """Hàm phụ: Đọc file Excel (.xlsx) và chuyển thành bảng Markdown"""
        excel_data = pd.read_excel(file_path, sheet_name=None)
        pages_text = []
        for sheet_name, df in excel_data.items():
            df = df.dropna(how='all').dropna(axis=1, how='all')
            if not df.empty:
                markdown_table = df.to_markdown(index=False)
                text = f"--- Dữ liệu từ Sheet: {sheet_name} ---\n{markdown_table}"
                pages_text.append({"text": text, "page": f"Sheet {sheet_name}"})
        return pages_text

    def process_pptx(self, file_path: str):
        """Hàm phụ: Đọc file PowerPoint (.pptx)"""
        prs = Presentation(file_path)
        pages_text = []
        for i, slide in enumerate(prs.slides):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    slide_text.append(shape.text)
            
            full_text = "\n".join([t for t in slide_text if t.strip()])
            if full_text:
                pages_text.append({"text": full_text, "page": i + 1})
        return pages_text

    def process_file(self, file_path: str, original_filename: str = None):
        """Đọc PDF/Docx/Xlsx/Pptx, dọn dẹp Text, bóc tách theo Cấu trúc (Heading) và gán Siêu dữ liệu phân cấp"""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Không tìm thấy file: {file_path}")

        print(f"[INFO] Dang doc va boc tach cau truc file: {file_path}")
        doc_display_name = original_filename or os.path.basename(file_path)
        
        file_ext = file_path.lower()
        if file_ext.endswith(".pdf"):
            pages_data = self.process_pdf(file_path)
        elif file_ext.endswith((".png", ".jpg", ".jpeg")):
            pages_data = self.process_image(file_path)
        elif file_ext.endswith(".docx"):
            pages_data = self.process_docx(file_path)
        elif file_ext.endswith(".xlsx"):
            pages_data = self.process_xlsx(file_path)
        elif file_ext.endswith(".pptx"):
            pages_data = self.process_pptx(file_path)
        else:
            raise ValueError("Định dạng file không được hỗ trợ (Chỉ nhận .pdf, .docx, .xlsx, .pptx, .png, .jpg, .jpeg)")

        chunks_with_metadata = []

        # Kiểm tra văn bản pháp luật (Chứa "Điều 1.", "Điều 2.")
        full_document_text = "\n".join([self.clean_text(d["text"]) for d in pages_data])
        is_legal_doc = bool(re.search(r'^Điều\s+\d+[\.:]?', full_document_text, re.MULTILINE | re.IGNORECASE))
        
        if is_legal_doc and file_ext in [".pdf", ".docx"]:
            print("[INFO] Phat hien van ban Phap luat -> Kich hoat LegalDocumentParser")
            parser = LegalDocumentParser({"title": doc_display_name, "source": doc_display_name})
            legal_docs = parser.parse(full_document_text)
            
            for idx, doc in enumerate(legal_docs):
                if len(doc.page_content.strip()) > 10 and not self.is_junk_chunk(doc.page_content):
                    meta = doc.metadata.copy()
                    meta.update({
                        "source": doc_display_name,
                        "filename": doc_display_name,
                        "section_title": meta.get("article_title") or "Dieu khoan",
                        "chunk_type": "legal",
                        "char_count": len(doc.page_content.strip())
                    })
                    chunks_with_metadata.append({
                        "content": doc.page_content.strip(),
                        "metadata": meta
                    })
            total = len(chunks_with_metadata)
            for idx, item in enumerate(chunks_with_metadata):
                item["metadata"]["chunk_index"] = idx
                item["metadata"]["total_chunks"] = total
            print(f"[INFO] Legal Chunking hoan tat: Tao ra {total} khoi theo cau truc Dieu/Khoan.")
            return chunks_with_metadata

        # Văn bản thông thường (Báo cáo, Luận văn, Tài liệu kỹ thuật, Word, Excel...)
        # Phân tích cấu trúc theo Heading/Chương mục
        print("[INFO] Kich hoat Structure-Aware Chunking (Phan tich cau truc theo Chuong muc & Tieu de)...")
        current_heading = "Tổng quan"
        raw_sections = []
        
        for data in pages_data:
            page_num = data["page"]
            text = data["text"]
            lines = text.split("\n")
            page_content_lines = []
            
            for line in lines:
                line_str = line.strip()
                if not line_str:
                    continue
                # Bỏ qua dòng rác số trang đơn độc ở chân/đầu trang (Page numbering)
                if re.match(r'^\d+$', line_str) or re.match(r'^page\s+\d+(\s+of\s+\d+)?$', line_str, re.IGNORECASE):
                    continue
                    
                heading_cand = self.extract_heading(line_str)
                if heading_cand:
                    if page_content_lines:
                        raw_sections.append((page_num, current_heading, "\n".join(page_content_lines)))
                        page_content_lines = []
                    current_heading = heading_cand
                else:
                    page_content_lines.append(line_str)
                    
            if page_content_lines:
                raw_sections.append((page_num, current_heading, "\n".join(page_content_lines)))

        for page_num, heading, text in raw_sections:
            clean_text = self.clean_text(text)
            if len(clean_text) < 40:
                continue
            sub_chunks = self.recursive_splitter.split_text(clean_text)
            chunk_type = self.classify_chunk_type(heading, clean_text)
            for sc in sub_chunks:
                sc_clean = sc.strip()
                if len(sc_clean) > 60 and not self.is_junk_chunk(sc_clean):
                    chunks_with_metadata.append({
                        "content": sc_clean,
                        "metadata": {
                            "source": doc_display_name,
                            "filename": doc_display_name,
                            "page": page_num,
                            "page_end": page_num,
                            "section_title": heading,
                            "chunk_type": chunk_type,
                            "char_count": len(sc_clean)
                        }
                    })

        total = len(chunks_with_metadata)
        for idx, item in enumerate(chunks_with_metadata):
            item["metadata"]["chunk_index"] = idx
            item["metadata"]["total_chunks"] = total
            
        print(f"[INFO] Structure-Aware Chunking hoan tat: Tao ra {total} khoi tri thuc co cau truc hoan chinh.")
        return chunks_with_metadata
